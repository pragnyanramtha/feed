#!/usr/bin/env python3
"""Hackathon pitch — 7 slides, exact judging template.

Edit the TEAM block below, then: python pitch/generate_pptx.py
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# ── paste your details here ──────────────────────────────────────────
PROJECT = "Campus Notice Desk"
THEME = "Notion Track"
TEAM_NAME = "Feed"
TEAM_ID = "TBD"
COLLEGE = "Malla Reddy Vishwavidyapeeth"
# ─────────────────────────────────────────────────────────────────────

OUT = Path(__file__).resolve().parent / "Campus-Notice-Desk.pptx"
W, H = Inches(13.333), Inches(7.5)

BG = RGBColor(0x12, 0x12, 0x12)
INK = RGBColor(0xF4, 0xF1, 0xEA)
MUTED = RGBColor(0x8A, 0x85, 0x7A)
ACCENT = RGBColor(0xFF, 0x5C, 0x39)
LIME = RGBColor(0xC4, 0xF5, 0x42)
CARD = RGBColor(0x1C, 0x1C, 0x1C)

TOTAL = 7


def _fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _run(p, text, *, size=18, bold=False, color=INK) -> None:
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def add_text(slide, l, t, w, h, lines, *, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    if isinstance(lines, str):
        lines = [lines]
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4)
        if isinstance(item, tuple):
            text, kw = item
            _run(p, text, size=kw.get("size", size), bold=kw.get("bold", bold), color=kw.get("color", color))
        else:
            _run(p, item, size=size, bold=bold, color=color)
    return box


def pill(slide, l, t, w, h, text, fill, color=BG):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    _fill(sh, fill)
    sh.adjustments[0] = 0.5
    tf = sh.text_frame
    tf.word_wrap = False
    tf.margin_top = Pt(7)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, text, size=12, bold=True, color=color)
    return sh


def card(slide, l, t, w, h):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    _fill(sh, CARD)
    sh.adjustments[0] = 0.08
    return sh


def slide(prs: Presentation):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    _fill(bg, BG)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), H)
    _fill(bar, ACCENT)
    return s


def footer(s, n: int) -> None:
    add_text(s, Inches(0.55), Inches(7.12), Inches(9), Inches(0.28), f"{PROJECT}  ·  {THEME}", size=10, color=MUTED)
    add_text(s, Inches(11.3), Inches(7.12), Inches(1.6), Inches(0.28), f"{n:02d}  /  {TOTAL:02d}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def kicker(s, text: str) -> None:
    add_text(s, Inches(0.7), Inches(0.28), Inches(12), Inches(0.32), text.upper(), size=12, bold=True, color=ACCENT)


def build() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # ── 1. Title ────────────────────────────────────────────────────
    s = slide(prs)
    kicker(s, THEME)
    add_text(s, Inches(0.7), Inches(1.35), Inches(12), Inches(1.2), PROJECT, size=48, bold=True)
    add_text(
        s,
        Inches(0.7),
        Inches(2.7),
        Inches(11.5),
        Inches(0.9),
        "WhatsApp is the pipe. Notion is the desk. Your code is the engine.",
        size=20,
        color=MUTED,
    )
    meta = [
        ("Theme", THEME),
        ("Team", TEAM_NAME),
        ("Team ID", TEAM_ID),
        ("College", COLLEGE),
    ]
    for i, (label, value) in enumerate(meta):
        x = Inches(0.7) + Inches((i % 4) * 3.05)
        y = Inches(4.35)
        card(s, x, y, Inches(2.9), Inches(1.85))
        add_text(s, x + Inches(0.22), y + Inches(0.28), Inches(2.5), Inches(0.35), label.upper(), size=11, bold=True, color=ACCENT)
        add_text(s, x + Inches(0.22), y + Inches(0.7), Inches(2.5), Inches(0.9), value, size=16, bold=True)
    footer(s, 1)

    # ── 2. Problem + Existing Gap ───────────────────────────────────
    s = slide(prs)
    kicker(s, "Problem + existing gap")
    add_text(s, Inches(0.7), Inches(0.7), Inches(12), Inches(0.7), "College life already runs on WhatsApp. WhatsApp is a chat, not a notice board.", size=22, bold=True)

    blocks = [
        ("The problem", "Circulars, exams, fests, “submit by 5pm” all land in the official group and die under stickers and forwards."),
        ("Who is affected", "Students who miss deadlines. CRs who retype the same message. Clubs and offices that have no filing system."),
        ("Current solutions", "ERP → WhatsApp blasts. Student apps that scrape the group. Notion templates you fill by hand. Zapier dumps."),
        ("The gap", "Nothing is a running service with Notion as the desk, a human approve-gate, a real send, and a Run Log written by code."),
    ]
    for i, (title, body) in enumerate(blocks):
        col, row = i % 2, i // 2
        x, y = Inches(0.55 + col * 6.3), Inches(1.65 + row * 2.5)
        card(s, x, y, Inches(6.05), Inches(2.3))
        add_text(s, x + Inches(0.3), y + Inches(0.25), Inches(5.45), Inches(0.4), title, size=16, bold=True, color=LIME)
        add_text(s, x + Inches(0.3), y + Inches(0.75), Inches(5.45), Inches(1.3), body, size=15, color=INK)
    footer(s, 2)

    # ── 3. Proposed Solution ────────────────────────────────────────
    s = slide(prs)
    kicker(s, "Proposed solution")
    add_text(s, Inches(0.7), Inches(0.7), Inches(12), Inches(0.85), "One shared notice desk. Same board for the whole class — categorized, dated, approved, then sent.", size=22, bold=True)

    add_text(s, Inches(0.7), Inches(1.7), Inches(12), Inches(0.35), "In simple terms", size=13, bold=True, color=ACCENT)
    add_text(
        s,
        Inches(0.7),
        Inches(2.05),
        Inches(12),
        Inches(0.85),
        "Forward the messy WhatsApp into our number. We file it in Notion. A human hits Approve. A clean notice goes out.",
        size=18,
    )

    trio = [
        ("How it solves it", "The CR stops rereading 200 chats. Students open Notion instead of scrolling. Deadlines sit on the card, not in the thread."),
        ("Better than a feed", "A labeled dump is a dashboard. We pause for approval, send outside Notion, and write a Run Log. That is the job, killed."),
        ("Better than clones", "Extracta / Chatnalyxer are student apps. ERPs push the other way. We keep WhatsApp as the front door and Notion as ops."),
    ]
    for i, (title, body) in enumerate(trio):
        x = Inches(0.55) + Inches(i * 4.15)
        card(s, x, Inches(3.15), Inches(3.95), Inches(3.2))
        add_text(s, x + Inches(0.25), Inches(3.4), Inches(3.45), Inches(0.55), title, size=16, bold=True, color=LIME)
        add_text(s, x + Inches(0.25), Inches(4.05), Inches(3.45), Inches(2.0), body, size=14)
    footer(s, 3)

    # ── 4. Architecture + Tech Stack ────────────────────────────────
    s = slide(prs)
    kicker(s, "Architecture + tech stack")
    add_text(s, Inches(0.7), Inches(0.68), Inches(12), Inches(0.4), "Input → Process → Output. Notion is the UI. This repo is the engine.", size=18, bold=True)

    flow = [
        ("Input", "WhatsApp Cloud API\nPOST /inbound\nforwarded circular"),
        ("Process", "Rules + Gemini parse\nDedupe fingerprint\nHold garbage / human"),
        ("Desk", "Notion Notices\nApprove / Reject\nOverride draft"),
        ("Output", "WhatsApp / email\nDeadline reminder\nRun Log by code"),
    ]
    for i, (title, body) in enumerate(flow):
        x = Inches(0.45) + Inches(i * 3.2)
        card(s, x, Inches(1.25), Inches(3.0), Inches(2.55))
        add_text(s, x + Inches(0.2), Inches(1.4), Inches(2.6), Inches(0.4), f"0{i+1}  {title}", size=14, bold=True, color=ACCENT)
        add_text(s, x + Inches(0.2), Inches(1.9), Inches(2.6), Inches(1.7), body, size=13, color=INK)
        if i < 3:
            add_text(s, x + Inches(2.85), Inches(2.15), Inches(0.4), Inches(0.4), "→", size=18, color=LIME)

    stack = [
        ("Frontend / UI", "Notion workspace\n(no React app)"),
        ("Backend", "Python · FastAPI\nVercel / Render"),
        ("AI", "If-statements first\nGemini / OpenAI"),
        ("DB + audit", "Notion DBs\nInbox · Notices · Log"),
        ("Action", "WA Cloud API\nResend email"),
        ("Triggers", "Webhooks + cron\n/health live"),
    ]
    for i, (title, body) in enumerate(stack):
        x = Inches(0.45) + Inches(i * 2.13)
        card(s, x, Inches(4.05), Inches(2.02), Inches(2.35))
        add_text(s, x + Inches(0.12), Inches(4.18), Inches(1.8), Inches(0.55), title, size=12, bold=True, color=LIME)
        add_text(s, x + Inches(0.12), Inches(4.75), Inches(1.8), Inches(1.4), body, size=12, color=INK)
    footer(s, 4)

    # ── 5. Key Features ─────────────────────────────────────────────
    s = slide(prs)
    kicker(s, "Key features")
    add_text(s, Inches(0.7), Inches(0.68), Inches(12), Inches(0.4), "Five. One job. Nothing decorative.", size=18, bold=True)

    feats = [
        ("01", "Runs without us", "Webhook or cron on a deployed host. Hitting a script during demo is not a service."),
        ("02", "Human approval in Notion", "Pending → Approve / Reject / edit draft. The send does not fire until a person decides."),
        ("03", "AI only where rules fail", "Hinglish circular → title, category, deadline, draft. “Good morning” is an if-statement."),
        ("04", "Real-world send", "Clean WhatsApp or email digest. Tomorrow-morning reminder for due notices. Not a chart."),
        ("05", "Proof + no silent loss", "Run Log written by the integration. Duplicates collapse. Garbage is held, not dropped."),
    ]
    for i, (num, title, body) in enumerate(feats):
        y = Inches(1.2) + Inches(i * 1.05)
        card(s, Inches(0.55), y, Inches(12.2), Inches(0.95))
        add_text(s, Inches(0.8), y + Inches(0.22), Inches(0.7), Inches(0.5), num, size=16, bold=True, color=ACCENT)
        add_text(s, Inches(1.6), y + Inches(0.12), Inches(10.8), Inches(0.35), title, size=16, bold=True, color=LIME)
        add_text(s, Inches(1.6), y + Inches(0.48), Inches(10.8), Inches(0.4), body, size=14, color=INK)
    footer(s, 5)

    # ── 6. Impact + Feasibility ─────────────────────────────────────
    s = slide(prs)
    kicker(s, "Impact + feasibility")
    add_text(s, Inches(0.7), Inches(0.7), Inches(12), Inches(0.45), "Built to leave behind. Free tiers. One class first.", size=20, bold=True)

    cells = [
        ("Who benefits", "Students who stop missing circulars. CRs who stop forwarding. Clubs and college offices with no ERP."),
        ("Real-world impact", "The weekly job of “read the group, decide, ping people” is gone. Notion is still usable if the engine is off."),
        ("Scalability", "One workspace per class or club. Same engine. Filter by audience later — not 2,000 personal feeds."),
        ("Feasibility", "Notion Education Plus. Meta test number. Gemini free tier. Vercel/Render free host. No GST on day one."),
    ]
    for i, (title, body) in enumerate(cells):
        col, row = i % 2, i // 2
        x, y = Inches(0.55 + col * 6.3), Inches(1.4 + row * 2.55)
        card(s, x, y, Inches(6.05), Inches(2.35))
        add_text(s, x + Inches(0.3), y + Inches(0.28), Inches(5.45), Inches(0.4), title, size=16, bold=True, color=LIME)
        add_text(s, x + Inches(0.3), y + Inches(0.8), Inches(5.45), Inches(1.3), body, size=15)
    footer(s, 6)

    # ── 7. Demo + Conclusion ────────────────────────────────────────
    s = slide(prs)
    kicker(s, "Demo + conclusion")
    add_text(s, Inches(0.7), Inches(0.68), Inches(12), Inches(0.45), "Live prototype. Do not pre-fill the Run Log.", size=20, bold=True)

    demo = [
        "POST /inbound or forward a real circular → Inbox + Notice card (by the integration)",
        "Send “gm” → ignored, logged, no card",
        "Approve in Notion → WhatsApp/email goes out → Run Log: sent",
        "Paste the same text again → duplicate, no second notice",
    ]
    for i, t in enumerate(demo):
        y = Inches(1.25) + Inches(i * 0.7)
        pill(s, Inches(0.7), y, Inches(0.5), Inches(0.42), str(i + 1), ACCENT, INK)
        add_text(s, Inches(1.4), y, Inches(11.2), Inches(0.5), t, size=16)

    card(s, Inches(0.55), Inches(4.15), Inches(12.2), Inches(2.2))
    add_text(s, Inches(0.85), Inches(4.35), Inches(11.6), Inches(0.35), "Why this should be selected", size=13, bold=True, color=ACCENT)
    add_text(
        s,
        Inches(0.85),
        Inches(4.75),
        Inches(11.6),
        Inches(0.55),
        "Delete the repo and it dies. Leave Notion open and a stranger can still run the job.",
        size=18,
        bold=True,
    )
    add_text(
        s,
        Inches(0.85),
        Inches(5.4),
        Inches(11.6),
        Inches(0.7),
        "One boring campus job, fully automated, with proof. That is the track.",
        size=16,
        color=MUTED,
    )
    footer(s, 7)

    prs.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
